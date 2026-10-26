"""
Parceria Comercial
Carrega a planilha de transações (aba "tratado"), lista os arquivos de comissão
de uma pasta, identifica o parceiro pelo nome do arquivo (via Arquivo - Parceiro.xlsx)
e preenche a aba "Transações" de cada arquivo com as transações do parceiro.

Dependências: pip install openpyxl
"""

import os
import queue
import re
import sys
import threading
import unicodedata
import webbrowser
import tkinter as tk
from datetime import date, datetime, time
from tkinter import filedialog, messagebox, ttk

from openpyxl import load_workbook

# ---------------------------------------------------------------------------
# Configuração
# ---------------------------------------------------------------------------
# Pasta do script (ou do .exe, se empacotado com PyInstaller)
if getattr(sys, "frozen", False):
    BASE_DIR = os.path.dirname(sys.executable)
else:
    BASE_DIR = os.path.dirname(os.path.abspath(__file__))

ARQUIVO_MAPEAMENTO = os.path.join(BASE_DIR, "Arquivo - Parceiro.xlsx")

VERSAO = "v1.0"
LINK_OUTROS_SISTEMAS = "http://192.168.16.65:10"


def recurso(nome):
    """Caminho de um arquivo embutido no .exe (ícone). No PyInstaller fica em sys._MEIPASS."""
    base = getattr(sys, "_MEIPASS", BASE_DIR)
    return os.path.join(base, nome)


# Planilha de transações
ABA_TRANSACOES_ORIGEM = "tratado"   # aba lida na planilha de transações
LINHA_CABECALHO = 2                 # cabeçalho na linha 2 (dados a partir da 3)
N_COLUNAS = 36                      # colunas A:AJ
COLUNA_PARCEIRO = "PARCEIRO"        # coluna usada para separar por parceiro

# Arquivos de comissão
ABA_TRANSACOES_DESTINO = "Transações"  # aba preenchida em cada arquivo
LINHA_INICIAL_DESTINO = 5              # dados gravados a partir da linha 5
POSICAO_PARCEIRO = 3                   # parte do nome (separada por "-") com o parceiro
EXTENSOES_ACEITAS = (".xlsx",)

# Tabela de transações: quantas linhas mostrar no máximo (a pesquisa filtra o resto)
LIMITE_EXIBICAO = 2000

# Caixas de seleção da coluna "Gerar"
MARCADO = "☑"
DESMARCADO = "☐"
BLOQUEADO = "—"  # linha com alerta: não pode ser gerada

# Paleta preto e branco (mesma do Envio de Comissões)
PRETO = "#000000"
GRAFITE = "#1a1a1a"
CINZA_ESCURO = "#333333"
CINZA = "#8a8a8a"
CINZA_CLARO = "#e6e6e6"
CINZA_LINHA = "#f5f5f5"
BRANCO = "#ffffff"
VERMELHO_CLARO = "#fbd5d5"
VERMELHO_TEXTO = "#7a1010"

FONTE = "Segoe UI"

# Estado do app
estado = {
    "transacoes": None,    # {arquivo, cabecalho, linhas, busca, por_parceiro, n_parceiros}
    "pasta": None,         # pasta dos arquivos de comissão
    "comissoes": [],       # [{arquivo, caminho, nome_busca, parceiro, erro, resultado, detalhe}]
    "desmarcados": set(),  # arquivos que o usuário tirou da geração
    "mapa": {},            # {ARQUIVO: PARCEIRO} lido do Arquivo - Parceiro.xlsx
    "mtime_mapa": None,    # última modificação do mapeamento lida
    "ocupado": False,      # True enquanto carrega ou gera em segundo plano
    "aba": "transacoes",   # aba visível: "transacoes" ou "comissoes"
    "filtradas": 0,        # registros encontrados na pesquisa
    "busca_agendada": None,
}

fila = queue.Queue()  # mensagens das tarefas em segundo plano para a tela


# ---------------------------------------------------------------------------
# Regras de negócio
# ---------------------------------------------------------------------------
def limpar_texto(texto):
    """Remove espaços duplicados e espaços no início/fim."""
    return re.sub(r"\s+", " ", str(texto)).strip()


def sem_acento(texto):
    """'Transações' -> 'transacoes' (minúsculo e sem acentos)."""
    normalizado = unicodedata.normalize("NFKD", str(texto))
    return "".join(c for c in normalizado if not unicodedata.combining(c)).lower()


def chave(texto):
    """Chave de comparação: maiúsculo, sem espaços extras."""
    return limpar_texto(texto).upper() if texto is not None else ""


def numero_br(n):
    """12345 -> '12.345'."""
    return f"{n:,}".replace(",", ".")


def formatar_valor(valor):
    """Texto mostrado na tabela de transações."""
    if valor is None:
        return ""
    if isinstance(valor, datetime):
        if valor.time() == time(0, 0):
            return valor.strftime("%d/%m/%Y")
        return valor.strftime("%d/%m/%Y %H:%M")
    if isinstance(valor, date):
        return valor.strftime("%d/%m/%Y")
    if isinstance(valor, bool):
        return "Sim" if valor else "Não"
    if isinstance(valor, float):
        if valor.is_integer() and abs(valor) < 1e15:
            return str(int(valor))
        texto = f"{valor:,.2f}" if round(valor, 2) == valor else f"{valor:,.6f}".rstrip("0")
        return texto.replace(",", "_").replace(".", ",").replace("_", ".")
    return str(valor)


def extrair_nome_busca(nome_arquivo):
    """3ª parte do nome (separando por '-'), em maiúsculo. '' se o nome não tiver partes suficientes."""
    partes = os.path.splitext(nome_arquivo)[0].upper().split("-")
    if len(partes) < POSICAO_PARCEIRO:
        return ""
    return partes[POSICAO_PARCEIRO - 1].strip()


def carregar_mapeamento(caminho):
    """Lê o Arquivo - Parceiro.xlsx e devolve {ARQUIVO: PARCEIRO}."""
    wb = load_workbook(caminho, read_only=True, data_only=True)
    try:
        ws = wb.active
        linhas = ws.iter_rows(values_only=True)
        cabecalho = [chave(c) for c in next(linhas, [])]
        if "ARQUIVO" not in cabecalho or "PARCEIRO" not in cabecalho:
            raise ValueError("O Excel precisa ter as colunas 'ARQUIVO' e 'PARCEIRO'.")
        i_arq = cabecalho.index("ARQUIVO")
        i_parc = cabecalho.index("PARCEIRO")

        mapa = {}
        for linha in linhas:
            if not linha or len(linha) <= max(i_arq, i_parc) or linha[i_arq] is None:
                continue
            arquivo = chave(linha[i_arq])
            parceiro = limpar_texto(linha[i_parc]) if linha[i_parc] is not None else ""
            if arquivo and arquivo not in mapa:  # vale a primeira ocorrência
                mapa[arquivo] = parceiro
        return mapa
    finally:
        wb.close()


def ler_transacoes(caminho, avisar_progresso):
    """
    Lê a aba 'tratado' (cabeçalho na linha 2, colunas A:AJ).
    Devolve {arquivo, cabecalho, linhas, busca, por_parceiro, n_parceiros}.
    """
    wb = load_workbook(caminho, read_only=True, data_only=True)
    try:
        aba = next((a for a in wb.sheetnames if sem_acento(a.strip()) == ABA_TRANSACOES_ORIGEM), None)
        if aba is None:
            raise ValueError(f"A planilha precisa ter a aba '{ABA_TRANSACOES_ORIGEM}'.")
        ws = wb[aba]
        total_estimado = max((ws.max_row or 0) - LINHA_CABECALHO, 0)

        iterador = ws.iter_rows(min_row=LINHA_CABECALHO, max_col=N_COLUNAS, values_only=True)
        bruto = list(next(iterador, ()))
        bruto += [None] * (N_COLUNAS - len(bruto))
        cabecalho = [
            limpar_texto(c) if c is not None and str(c).strip() else f"Coluna {i + 1}"
            for i, c in enumerate(bruto[:N_COLUNAS])
        ]

        i_parceiro = next((i for i, c in enumerate(cabecalho) if chave(c) == COLUNA_PARCEIRO), None)
        if i_parceiro is None:
            raise ValueError(f"A aba '{aba}' precisa ter a coluna '{COLUNA_PARCEIRO}' (cabeçalho na linha {LINHA_CABECALHO}).")

        linhas, busca, por_parceiro = [], [], {}
        for n, linha in enumerate(iterador, start=1):
            linha = tuple(linha[:N_COLUNAS]) + (None,) * (N_COLUNAS - len(linha))
            if all(v is None or v == "" for v in linha):
                continue  # linha totalmente vazia
            linhas.append(linha)
            busca.append(sem_acento(" ".join(formatar_valor(v) for v in linha)))
            por_parceiro.setdefault(chave(linha[i_parceiro]), []).append(linha)
            if n % 2000 == 0:
                avisar_progresso(n, total_estimado)

        return {
            "arquivo": caminho,
            "cabecalho": cabecalho,
            "linhas": linhas,
            "busca": busca,
            "por_parceiro": por_parceiro,
            "n_parceiros": len([p for p in por_parceiro if p]),
        }
    finally:
        wb.close()


def listar_arquivos(pasta):
    """Arquivos .xlsx da pasta (sem subpastas), ignorando ocultos e temporários do Office."""
    return sorted(
        f for f in os.listdir(pasta)
        if os.path.isfile(os.path.join(pasta, f))
        and f.lower().endswith(EXTENSOES_ACEITAS)
        and not f.startswith(("~$", "."))
    )


def montar_comissoes(pasta, mapa):
    """Junta arquivo + nome extraído + parceiro do mapeamento."""
    linhas = []
    for nome in listar_arquivos(pasta):
        nome_busca = extrair_nome_busca(nome)
        parceiro = mapa.get(nome_busca) if nome_busca else None
        if not nome_busca:
            erro = f"nome fora do padrão (parceiro é a {POSICAO_PARCEIRO}ª parte, separada por '-')"
        elif parceiro is None:
            erro = f"'{nome_busca}' não está na coluna ARQUIVO do Arquivo - Parceiro"
        else:
            erro = ""
        linhas.append({
            "arquivo": nome,
            "caminho": os.path.normpath(os.path.join(pasta, nome)),
            "nome_busca": nome_busca,
            "parceiro": parceiro or "",
            "erro": erro,
            "resultado": None,  # None, "ok" ou "erro" (depois de gerar)
            "detalhe": "",
        })
    return linhas


def transacoes_do_parceiro(linha):
    """Linhas de transação do parceiro do arquivo (lista vazia se não houver)."""
    if estado["transacoes"] is None:
        return []
    return estado["transacoes"]["por_parceiro"].get(chave(linha["parceiro"]), [])


def pode_gerar(linha):
    """Arquivo sem alerta: parceiro identificado no mapeamento."""
    return not linha["erro"]


def vai_gerar(linha):
    return pode_gerar(linha) and linha["arquivo"] not in estado["desmarcados"]


def preencher_arquivo(caminho, transacoes):
    """
    Limpa a aba 'Transações' (a partir da linha 5, colunas A:AJ) e grava as transações.
    Salva primeiro numa cópia temporária e só depois troca o arquivo original,
    para não corromper a planilha se algo der errado no meio.
    """
    wb = load_workbook(caminho, keep_links=True)
    try:
        aba = next(
            (a for a in wb.sheetnames if sem_acento(a.strip()) == sem_acento(ABA_TRANSACOES_DESTINO)),
            None,
        )
        if aba is None:
            raise ValueError(f"aba '{ABA_TRANSACOES_DESTINO}' não encontrada")
        ws = wb[aba]

        for row in ws.iter_rows(
            min_row=LINHA_INICIAL_DESTINO, max_row=max(ws.max_row, LINHA_INICIAL_DESTINO),
            min_col=1, max_col=N_COLUNAS,
        ):
            for cell in row:
                cell.value = None

        for numero_linha, valores in enumerate(transacoes, start=LINHA_INICIAL_DESTINO):
            for numero_coluna, valor in enumerate(valores, start=1):
                ws.cell(row=numero_linha, column=numero_coluna, value=valor)

        pasta, nome = os.path.split(caminho)
        temporario = os.path.join(pasta, f"~tmp_{nome}")
        wb.save(temporario)
    finally:
        wb.close()

    try:
        os.replace(temporario, caminho)
    except Exception:
        try:
            os.remove(temporario)
        except OSError:
            pass
        raise


def mensagem_de_erro(erro):
    """Texto curto e compreensível para a coluna 'Situação'."""
    if isinstance(erro, PermissionError):
        return "arquivo aberto no Excel ou sem permissão — feche e gere de novo"
    if isinstance(erro, AttributeError) and "read-only" in str(erro):
        return "a aba Transações tem células mescladas na área de dados"
    return str(erro) or erro.__class__.__name__


# ---------------------------------------------------------------------------
# Componentes visuais
# ---------------------------------------------------------------------------
def criar_botao(pai, texto, comando, primario=True):
    """Botão plano. Primário = preto; secundário = branco com borda preta."""
    fundo, letra = (PRETO, BRANCO) if primario else (BRANCO, PRETO)
    fundo_hover = CINZA_ESCURO if primario else CINZA_CLARO

    botao = tk.Label(
        pai,
        text=texto,
        font=(FONTE, 11, "bold"),
        fg=letra,
        bg=fundo,
        padx=24,
        pady=8,
        cursor="hand2",
        highlightthickness=2,
        highlightbackground=PRETO,
    )
    botao.bind("<Button-1>", lambda e: comando())
    botao.bind("<Enter>", lambda e: botao.config(bg=fundo_hover))
    botao.bind("<Leave>", lambda e: botao.config(bg=fundo))
    return botao


def criar_link(pai, texto, comando):
    """Texto clicável sublinhado."""
    link = tk.Label(pai, text=texto, font=(FONTE, 9, "underline"), fg=PRETO, bg=BRANCO, cursor="hand2")
    link.bind("<Button-1>", lambda e: comando())
    return link


def criar_campo(pai):
    """Campo de uma linha com borda fina."""
    return tk.Entry(
        pai,
        font=(FONTE, 10),
        fg=GRAFITE,
        bg=BRANCO,
        relief="flat",
        highlightthickness=1,
        highlightbackground=CINZA_CLARO,
        highlightcolor=PRETO,
    )


def configurar_estilos(janela):
    style = ttk.Style(janela)
    style.theme_use("clam")

    # Tabelas
    style.configure(
        "Treeview",
        background=BRANCO,
        fieldbackground=BRANCO,
        foreground=GRAFITE,
        rowheight=32,
        font=(FONTE, 10),
        borderwidth=0,
    )
    style.configure(
        "Treeview.Heading",
        background=PRETO,
        foreground=BRANCO,
        font=(FONTE, 10, "bold"),
        relief="flat",
        padding=(10, 8),
    )
    style.map("Treeview.Heading", background=[("active", CINZA_ESCURO)])
    style.map(
        "Treeview",
        background=[("selected", CINZA_ESCURO)],
        foreground=[("selected", BRANCO)],
    )
    style.layout("Treeview", [("Treeview.treearea", {"sticky": "nswe"})])

    # Barras de rolagem finas, sem setas
    for orientacao in ("Vertical", "Horizontal"):
        nome = f"Fina.{orientacao}.TScrollbar"
        style.layout(nome, [
            (f"{orientacao}.Scrollbar.trough", {
                "sticky": "nswe",
                "children": [(f"{orientacao}.Scrollbar.thumb", {"expand": "1", "sticky": "nswe"})],
            }),
        ])
        style.configure(
            nome,
            background=CINZA,
            troughcolor=CINZA_LINHA,
            bordercolor=CINZA_LINHA,
            lightcolor=CINZA,
            darkcolor=CINZA,
            gripcount=0,
            arrowsize=10,
            relief="flat",
        )
        style.map(
            nome,
            background=[("pressed", PRETO), ("active", CINZA_ESCURO)],
            lightcolor=[("pressed", PRETO), ("active", CINZA_ESCURO)],
            darkcolor=[("pressed", PRETO), ("active", CINZA_ESCURO)],
        )

    # Barra de progresso fina e preta
    style.configure(
        "Preta.Horizontal.TProgressbar",
        background=PRETO,
        troughcolor=CINZA_CLARO,
        bordercolor=CINZA_CLARO,
        lightcolor=PRETO,
        darkcolor=PRETO,
        thickness=6,
    )


def criar_tabela(pai, colunas):
    """Treeview com barras de rolagem finas dentro de uma moldura de borda clara."""
    area = tk.Frame(pai, bg=BRANCO, highlightbackground=CINZA_CLARO, highlightthickness=1)
    tabela = ttk.Treeview(area, columns=colunas, show="headings", selectmode="browse")
    scroll_y = ttk.Scrollbar(area, orient="vertical", command=tabela.yview, style="Fina.Vertical.TScrollbar")
    scroll_x = ttk.Scrollbar(area, orient="horizontal", command=tabela.xview, style="Fina.Horizontal.TScrollbar")
    tabela.configure(yscrollcommand=scroll_y.set, xscrollcommand=scroll_x.set)

    tabela.grid(row=0, column=0, sticky="nswe")
    scroll_y.grid(row=0, column=1, sticky="ns")
    scroll_x.grid(row=1, column=0, sticky="we")
    area.rowconfigure(0, weight=1)
    area.columnconfigure(0, weight=1)

    tabela.tag_configure("par", background=BRANCO)
    tabela.tag_configure("impar", background=CINZA_LINHA)
    tabela.tag_configure("alerta", foreground=PRETO, font=(FONTE, 10, "bold"))
    tabela.tag_configure("erro", background=VERMELHO_CLARO, foreground=VERMELHO_TEXTO)

    # Shift + rodinha do mouse rola para os lados
    tabela.bind("<Shift-MouseWheel>", lambda e: tabela.xview_scroll(-1 * (e.delta // 120) * 3, "units"))
    return area, tabela


# ---------------------------------------------------------------------------
# Abas
# ---------------------------------------------------------------------------
def mostrar_aba(nome):
    estado["aba"] = nome
    for chave_aba, (rotulo, barra) in abas.items():
        ativa = chave_aba == nome
        rotulo.config(fg=PRETO if ativa else CINZA)
        barra.config(bg=PRETO if ativa else BRANCO)

    if nome == "transacoes":
        painel_comissoes.pack_forget()
        painel_transacoes.pack(fill="both", expand=True, padx=32)
    else:
        painel_transacoes.pack_forget()
        painel_comissoes.pack(fill="both", expand=True, padx=32)
    atualizar_status()


def criar_aba(pai, texto, nome):
    caixa = tk.Frame(pai, bg=BRANCO, cursor="hand2")
    rotulo = tk.Label(caixa, text=texto, font=(FONTE, 11, "bold"), fg=CINZA, bg=BRANCO, cursor="hand2")
    rotulo.pack(padx=2, pady=(0, 6))
    barra = tk.Frame(caixa, bg=BRANCO, height=3)
    barra.pack(fill="x")
    for widget in (caixa, rotulo):
        widget.bind("<Button-1>", lambda e: mostrar_aba(nome))
    abas[nome] = (rotulo, barra)
    return caixa


# ---------------------------------------------------------------------------
# Status, progresso e botões
# ---------------------------------------------------------------------------
def atualizar_status():
    """Texto do rodapé, conforme a aba visível."""
    if estado["ocupado"]:
        return  # durante o processamento o rodapé mostra o progresso

    if estado["aba"] == "transacoes":
        dados = estado["transacoes"]
        if dados is None:
            lbl_status.config(text="Carregue a planilha de transações")
            return
        partes = [
            f"{numero_br(len(dados['linhas']))} linhas",
            f"{len(dados['cabecalho'])} colunas",
            f"{dados['n_parceiros']} parceiros",
        ]
        if campo_busca_tem_texto():
            partes.append(f"{numero_br(estado['filtradas'])} encontrados")
        if estado["filtradas"] > LIMITE_EXIBICAO:
            partes.append(f"mostrando os primeiros {numero_br(LIMITE_EXIBICAO)}")
        lbl_status.config(text="  ·  ".join(partes))
    else:
        linhas = estado["comissoes"]
        if not estado["pasta"]:
            lbl_status.config(text="Busque a pasta dos arquivos de comissão")
            return
        partes = [f"{len(linhas)} arquivo(s)", f"{sum(1 for l in linhas if vai_gerar(l))} para gerar"]
        com_erro = sum(1 for l in linhas if l["erro"])
        if com_erro:
            partes.append(f"{com_erro} sem parceiro")
        if estado["transacoes"] is not None:
            sem_transacoes = sum(1 for l in linhas if pode_gerar(l) and not transacoes_do_parceiro(l))
            if sem_transacoes:
                partes.append(f"{sem_transacoes} sem transações")
        gerados = sum(1 for l in linhas if l["resultado"] == "ok")
        falhas = sum(1 for l in linhas if l["resultado"] == "erro")
        if gerados:
            partes.append(f"{gerados} gerado(s)")
        if falhas:
            partes.append(f"{falhas} com erro")
        lbl_status.config(text="  ·  ".join(partes))


def iniciar_ocupado(texto):
    estado["ocupado"] = True
    janela.config(cursor="watch")
    lbl_status.config(text=texto)
    barra_progresso.config(mode="indeterminate", value=0)
    barra_progresso.pack(side="left", padx=(0, 12), before=lbl_status)
    barra_progresso.start(12)


def progresso(texto, atual=None, total=None):
    lbl_status.config(text=texto)
    if atual is not None and total:
        if str(barra_progresso.cget("mode")) != "determinate":
            barra_progresso.stop()
            barra_progresso.config(mode="determinate", maximum=total)
        barra_progresso.config(value=min(atual, total))


def finalizar_ocupado():
    estado["ocupado"] = False
    janela.config(cursor="")
    barra_progresso.stop()
    barra_progresso.pack_forget()
    atualizar_status()


def atualizar_botao_gerar():
    """Mostra 'Gerar planilhas' só quando há transações carregadas e pasta escolhida."""
    pronto = estado["transacoes"] is not None and bool(estado["pasta"])
    if pronto and not btn_gerar.winfo_ismapped():
        btn_gerar.pack(side="right", padx=(0, 12), after=btn_pasta)  # à esquerda de "Buscar pasta"
    elif not pronto and btn_gerar.winfo_ismapped():
        btn_gerar.pack_forget()


def verificar_fila():
    """Recebe as mensagens das tarefas em segundo plano (a tela só é mexida aqui)."""
    try:
        while True:
            tipo, dados = fila.get_nowait()
            TRATADORES[tipo](dados)
    except queue.Empty:
        pass
    janela.after(80, verificar_fila)


# ---------------------------------------------------------------------------
# Transações
# ---------------------------------------------------------------------------
def carregar_transacoes():
    if estado["ocupado"]:
        return
    caminho = filedialog.askopenfilename(
        title="Selecione a planilha de transações",
        filetypes=[("Planilhas do Excel", "*.xlsx *.xlsm"), ("Todos os arquivos", "*.*")],
    )
    if not caminho:
        return

    mostrar_aba("transacoes")
    iniciar_ocupado("Carregando planilha...")

    def tarefa():
        try:
            def avisar(n, total):
                fila.put(("transacoes_progresso", (n, total)))
            fila.put(("transacoes_ok", ler_transacoes(caminho, avisar)))
        except Exception as erro:
            fila.put(("transacoes_erro", erro))

    threading.Thread(target=tarefa, daemon=True).start()


def ao_progresso_transacoes(dados):
    n, total = dados
    texto = f"Carregando planilha...  {numero_br(n)}"
    if total:
        texto += f" de ~{numero_br(total)} linhas"
    progresso(texto, n, total)


def ao_carregar_transacoes(dados):
    estado["transacoes"] = dados
    finalizar_ocupado()
    lbl_arquivo.config(text=os.path.normpath(dados["arquivo"]), fg=CINZA_ESCURO)
    montar_colunas_transacoes(dados)
    campo_busca.delete(0, "end")
    busca_sair()  # volta o texto de exemplo
    filtrar_transacoes()
    if not campo_busca.winfo_ismapped():
        campo_busca.pack(side="right", ipady=5)
        lbl_lupa.pack(side="right", padx=(0, 6))
    atualizar_botao_gerar()
    if estado["comissoes"]:
        for linha in estado["comissoes"]:
            linha["resultado"], linha["detalhe"] = None, ""
        preencher_tabela_comissoes()


def ao_erro_transacoes(erro):
    finalizar_ocupado()
    messagebox.showerror("Erro ao carregar a planilha", str(erro))


def montar_colunas_transacoes(dados):
    """Cria as colunas da tabela com largura conforme o conteúdo."""
    colunas = [f"c{i}" for i in range(len(dados["cabecalho"]))]
    tabela_transacoes.delete(*tabela_transacoes.get_children())
    tabela_transacoes.configure(columns=colunas)
    amostra = dados["linhas"][:200]
    for i, (coluna, titulo) in enumerate(zip(colunas, dados["cabecalho"])):
        maior = max([len(titulo)] + [len(formatar_valor(l[i])) for l in amostra])
        largura = min(max(maior * 8 + 24, 90), 320)
        tabela_transacoes.heading(coluna, text=titulo, anchor="w")
        tabela_transacoes.column(coluna, width=largura, minwidth=60, stretch=False, anchor="w")


def campo_busca_tem_texto():
    return bool(campo_busca.get().strip()) and campo_busca.cget("fg") != CINZA


def agendar_busca(evento=None):
    """Espera o usuário parar de digitar antes de filtrar."""
    if estado["busca_agendada"]:
        janela.after_cancel(estado["busca_agendada"])
    estado["busca_agendada"] = janela.after(250, filtrar_transacoes)


def filtrar_transacoes():
    estado["busca_agendada"] = None
    dados = estado["transacoes"]
    if dados is None:
        return
    termo = sem_acento(campo_busca.get().strip()) if campo_busca_tem_texto() else ""

    if termo:
        indices = [i for i, texto in enumerate(dados["busca"]) if termo in texto]
    else:
        indices = range(len(dados["linhas"]))
    estado["filtradas"] = len(indices)

    tabela_transacoes.delete(*tabela_transacoes.get_children())
    for posicao, i in enumerate(indices):
        if posicao >= LIMITE_EXIBICAO:
            break
        tabela_transacoes.insert(
            "", "end",
            values=[formatar_valor(v) for v in dados["linhas"][i]],
            tags=("par" if posicao % 2 == 0 else "impar",),
        )
    atualizar_status()


# Texto de exemplo dentro do campo de pesquisa
TEXTO_BUSCA = "Pesquisar em qualquer coluna"


def busca_entrar(evento=None):
    if campo_busca.cget("fg") == CINZA:
        campo_busca.delete(0, "end")
        campo_busca.config(fg=GRAFITE)


def busca_sair(evento=None):
    if not campo_busca.get():
        campo_busca.insert(0, TEXTO_BUSCA)
        campo_busca.config(fg=CINZA)


# ---------------------------------------------------------------------------
# Arquivos de comissão
# ---------------------------------------------------------------------------
def ler_mapeamento():
    """Lê o Arquivo - Parceiro.xlsx mostrando o erro na tela, se houver. None se falhar."""
    try:
        mapa = carregar_mapeamento(ARQUIVO_MAPEAMENTO)
    except FileNotFoundError:
        messagebox.showerror("Excel não encontrado", f"Não encontrei o arquivo:\n{ARQUIVO_MAPEAMENTO}")
        return None
    except Exception as erro:
        messagebox.showerror("Erro ao ler o Arquivo - Parceiro", str(erro))
        return None
    estado["mtime_mapa"] = os.path.getmtime(ARQUIVO_MAPEAMENTO)
    estado["mapa"] = mapa
    return mapa


def buscar_pasta():
    if estado["ocupado"]:
        return
    pasta = filedialog.askdirectory(title="Selecione a pasta dos arquivos de comissão")
    if not pasta:
        return
    mapa = ler_mapeamento()
    if mapa is None:
        return

    estado["pasta"] = pasta
    estado["desmarcados"] = set()  # pasta nova: todos marcados
    lbl_pasta.config(text=os.path.normpath(pasta), fg=CINZA_ESCURO)
    estado["comissoes"] = montar_comissoes(pasta, mapa)
    preencher_tabela_comissoes()
    mostrar_aba("comissoes")
    atualizar_botao_gerar()


def recarregar_se_mapeamento_mudou(evento=None):
    """Ao voltar para o app, relê o Arquivo - Parceiro se ele foi salvo nesse meio-tempo."""
    if estado["ocupado"] or not estado["pasta"] or not os.path.exists(ARQUIVO_MAPEAMENTO):
        return
    mtime = os.path.getmtime(ARQUIVO_MAPEAMENTO)
    if mtime == estado["mtime_mapa"]:
        return
    estado["mtime_mapa"] = mtime  # marca antes, para um erro não reabrir o aviso a cada clique
    mapa = ler_mapeamento()
    if mapa is not None:
        estado["comissoes"] = montar_comissoes(estado["pasta"], mapa)
        preencher_tabela_comissoes()


def abrir_mapeamento():
    """Abre o Arquivo - Parceiro.xlsx. Ao salvar e voltar, a tabela se atualiza sozinha."""
    if not os.path.exists(ARQUIVO_MAPEAMENTO):
        messagebox.showerror("Excel não encontrado", f"Não encontrei o arquivo:\n{ARQUIVO_MAPEAMENTO}")
        return
    abrir_no_sistema(ARQUIVO_MAPEAMENTO)


def abrir_no_sistema(caminho):
    try:
        os.startfile(caminho)  # só existe no Windows
    except AttributeError:
        messagebox.showinfo("Abrir arquivo", f"Abra manualmente:\n{caminho}")
    except OSError as erro:
        messagebox.showerror("Erro ao abrir", str(erro))


def caixa_da_linha(linha):
    if not pode_gerar(linha):
        return BLOQUEADO
    return DESMARCADO if linha["arquivo"] in estado["desmarcados"] else MARCADO


def situacao_da_linha(linha):
    """Texto da coluna 'Situação' e se a linha merece destaque (negrito)."""
    if linha["erro"]:
        return f"⚠ {linha['erro']}", False
    if linha["resultado"] == "ok":
        return f"✔ {linha['detalhe']}", False
    if linha["resultado"] == "erro":
        return f"✖ {linha['detalhe']}", False
    if linha["resultado"] == "processando":
        return "Processando...", True
    if estado["transacoes"] is None:
        return "Aguardando planilha de transações", False
    if not transacoes_do_parceiro(linha):
        return "⚠ sem transações — a aba Transações ficará vazia", True
    return "Pronto para gerar", False


def valores_da_linha(linha):
    if estado["transacoes"] is None or linha["erro"]:
        quantidade = "—"
    else:
        quantidade = numero_br(len(transacoes_do_parceiro(linha)))
    return (
        caixa_da_linha(linha),
        linha["arquivo"],
        linha["nome_busca"] or "—",
        linha["parceiro"] or "—",
        quantidade,
        situacao_da_linha(linha)[0],
    )


def tags_da_linha(indice, linha):
    if linha["erro"] or linha["resultado"] == "erro":
        return ["erro"]
    tags = ["par" if indice % 2 == 0 else "impar"]
    if situacao_da_linha(linha)[1]:
        tags.append("alerta")
    return tags


def preencher_tabela_comissoes():
    tabela_comissoes.delete(*tabela_comissoes.get_children())
    for i, linha in enumerate(estado["comissoes"]):
        tabela_comissoes.insert(
            "", "end", iid=linha["arquivo"], values=valores_da_linha(linha), tags=tags_da_linha(i, linha)
        )
    atualizar_status()


def atualizar_linha_comissao(linha):
    i = estado["comissoes"].index(linha)
    tabela_comissoes.item(linha["arquivo"], values=valores_da_linha(linha), tags=tags_da_linha(i, linha))


def linha_por_arquivo(arquivo):
    return next((l for l in estado["comissoes"] if l["arquivo"] == arquivo), None)


def clicar_na_tabela_comissoes(evento):
    """Clique na coluna 'Gerar' marca/desmarca o arquivo."""
    if estado["ocupado"]:
        return
    if tabela_comissoes.identify_region(evento.x, evento.y) != "cell":
        return
    if tabela_comissoes.identify_column(evento.x) != "#1":
        return
    linha = linha_por_arquivo(tabela_comissoes.identify_row(evento.y))
    if linha is None or not pode_gerar(linha):
        return
    estado["desmarcados"] ^= {linha["arquivo"]}
    tabela_comissoes.set(linha["arquivo"], "gerar", caixa_da_linha(linha))
    atualizar_status()


def duplo_clique_comissoes(evento):
    """Duplo clique em uma linha (fora da coluna 'Gerar') abre o arquivo no Excel."""
    if tabela_comissoes.identify_region(evento.x, evento.y) != "cell":
        return
    if tabela_comissoes.identify_column(evento.x) == "#1":
        return
    linha = linha_por_arquivo(tabela_comissoes.identify_row(evento.y))
    if linha is not None:
        abrir_no_sistema(linha["caminho"])


def alternar_todos():
    """Clique no título 'Gerar': se todos estão marcados, desmarca todos; senão, marca todos."""
    if estado["ocupado"]:
        return
    aptas = [l for l in estado["comissoes"] if pode_gerar(l)]
    if not aptas:
        return
    if all(vai_gerar(l) for l in aptas):
        estado["desmarcados"] = {l["arquivo"] for l in aptas}
    else:
        estado["desmarcados"] = set()
    for linha in aptas:
        tabela_comissoes.set(linha["arquivo"], "gerar", caixa_da_linha(linha))
    atualizar_status()


# ---------------------------------------------------------------------------
# Gerar planilhas
# ---------------------------------------------------------------------------
def gerar_planilhas():
    if estado["ocupado"]:
        return
    linhas = estado["comissoes"]
    aptas = [l for l in linhas if vai_gerar(l)]
    if not aptas:
        messagebox.showinfo("Nada a fazer", "Nenhum arquivo marcado e com parceiro identificado.")
        return

    sem_transacoes = sum(1 for l in aptas if not transacoes_do_parceiro(l))
    com_erro = sum(1 for l in linhas if not pode_gerar(l))
    desmarcadas = sum(1 for l in linhas if pode_gerar(l) and not vai_gerar(l))

    pergunta = (
        f"Gerar {len(aptas)} planilha(s)?\n\n"
        f"A aba '{ABA_TRANSACOES_DESTINO}' de cada arquivo será substituída "
        f"(a partir da linha {LINHA_INICIAL_DESTINO})."
    )
    if sem_transacoes:
        pergunta += f"\n\n{sem_transacoes} arquivo(s) ficarão com a aba vazia (parceiro sem transações)."
    ignorados = []
    if com_erro:
        ignorados.append(f"{com_erro} sem parceiro identificado")
    if desmarcadas:
        ignorados.append(f"{desmarcadas} desmarcado(s)")
    if ignorados:
        pergunta += "\n\nSerão ignorados: " + " e ".join(ignorados) + "."
    if not messagebox.askyesno("Confirmar", pergunta):
        return

    for linha in linhas:
        linha["resultado"], linha["detalhe"] = None, ""
    preencher_tabela_comissoes()

    trabalhos = [(l["arquivo"], l["caminho"], transacoes_do_parceiro(l)) for l in aptas]
    iniciar_ocupado(f"Gerando 0/{len(trabalhos)}...")
    progresso(f"Gerando 0/{len(trabalhos)}...", 0, len(trabalhos))

    def tarefa():
        for indice, (arquivo, caminho, transacoes) in enumerate(trabalhos):
            fila.put(("gerar_inicio", (arquivo, indice, len(trabalhos))))
            try:
                preencher_arquivo(caminho, transacoes)
                fila.put(("gerar_arquivo", (arquivo, "ok", f"{numero_br(len(transacoes))} transações gravadas")))
            except Exception as erro:
                fila.put(("gerar_arquivo", (arquivo, "erro", mensagem_de_erro(erro))))
        fila.put(("gerar_fim", None))

    threading.Thread(target=tarefa, daemon=True).start()


def ao_iniciar_arquivo(dados):
    arquivo, indice, total = dados
    progresso(f"Gerando {indice + 1}/{total}:  {arquivo}", indice, total)
    linha = linha_por_arquivo(arquivo)
    if linha is not None:
        linha["resultado"] = "processando"
        atualizar_linha_comissao(linha)
        tabela_comissoes.see(arquivo)


def ao_terminar_arquivo(dados):
    arquivo, resultado, detalhe = dados
    linha = linha_por_arquivo(arquivo)
    if linha is not None:
        linha["resultado"], linha["detalhe"] = resultado, detalhe
        atualizar_linha_comissao(linha)
    feitos = sum(1 for l in estado["comissoes"] if l["resultado"] in ("ok", "erro"))
    barra_progresso.config(value=feitos)


def ao_terminar_geracao(_):
    finalizar_ocupado()
    ok = [l for l in estado["comissoes"] if l["resultado"] == "ok"]
    erros = [l for l in estado["comissoes"] if l["resultado"] == "erro"]
    resumo = f"{len(ok)} planilha(s) gerada(s)."
    if erros:
        resumo += "\n\nErros:\n" + "\n".join(f"{l['arquivo']}: {l['detalhe']}" for l in erros)
        messagebox.showwarning("Concluído com erros", resumo)
    else:
        messagebox.showinfo("Concluído", resumo)


# Mensagem da fila -> função que atualiza a tela
TRATADORES = {
    "transacoes_progresso": ao_progresso_transacoes,
    "transacoes_ok": ao_carregar_transacoes,
    "transacoes_erro": ao_erro_transacoes,
    "gerar_inicio": ao_iniciar_arquivo,
    "gerar_arquivo": ao_terminar_arquivo,
    "gerar_fim": ao_terminar_geracao,
}


def fechar_janela():
    if estado["ocupado"]:
        if not messagebox.askyesno(
            "Processando",
            "Ainda há um processamento em andamento.\nFechar mesmo assim? O arquivo atual pode ficar incompleto.",
        ):
            return
    janela.destroy()


# ---------------------------------------------------------------------------
# Janela principal
# ---------------------------------------------------------------------------
janela = tk.Tk()
janela.title(f"Parceria Comercial OWN  —  {VERSAO}")

try:
    janela.iconbitmap(recurso("logo.ico"))
except Exception:
    try:
        janela._icone = tk.PhotoImage(file=recurso("logo.png"))
        janela.iconphoto(True, janela._icone)
    except Exception:
        pass  # sem ícone não impede o app de abrir
janela.geometry("1280x680")
janela.minsize(900, 500)
janela.configure(bg=BRANCO)
configurar_estilos(janela)

abas = {}

# --- Cabeçalho ---
topo = tk.Frame(janela, bg=BRANCO)
topo.pack(fill="x", padx=32, pady=(12, 0))

linha_titulo = tk.Frame(topo, bg=BRANCO)
linha_titulo.pack(fill="x")
tk.Label(
    linha_titulo, text="Parceria Comercial", font=(FONTE, 24, "bold"), fg=PRETO, bg=BRANCO
).pack(side="left")
tk.Label(
    linha_titulo, text=date.today().strftime("%d/%m/%Y"), font=(FONTE, 10), fg=PRETO, bg=BRANCO
).pack(side="right", anchor="n", pady=(6, 0))

# Abas logo acima da linha divisória
linha_abas = tk.Frame(topo, bg=BRANCO)
linha_abas.pack(fill="x", pady=(10, 0))
criar_aba(linha_abas, "Transações", "transacoes").pack(side="left")
criar_aba(linha_abas, "Arquivos de comissão", "comissoes").pack(side="left", padx=(24, 0))

# Links à direita, na altura das abas
criar_link(linha_abas, "Ver outros sistemas", lambda: webbrowser.open(LINK_OUTROS_SISTEMAS)).pack(
    side="right", anchor="n", pady=(4, 0)
)
criar_link(linha_abas, "Abrir planilha de parceiros", abrir_mapeamento).pack(
    side="right", anchor="n", pady=(4, 0), padx=(0, 16)
)

tk.Frame(janela, bg=PRETO, height=2).pack(fill="x", padx=32, pady=(0, 12))

# --- Rodapé (criado antes dos painéis para ficar sempre visível) ---
rodape = tk.Frame(janela, bg=BRANCO)
rodape.pack(side="bottom", fill="x", padx=32, pady=20)

# Botões primeiro, para nunca serem espremidos pelo texto de status
btn_transacoes = criar_botao(rodape, "Carregar transações", carregar_transacoes, primario=True)
btn_transacoes.pack(side="right")
btn_pasta = criar_botao(rodape, "Buscar pasta", buscar_pasta, primario=False)
btn_pasta.pack(side="right", padx=(0, 12))
# Aparece só quando há transações e pasta
btn_gerar = criar_botao(rodape, "Gerar planilhas", gerar_planilhas, primario=False)

lbl_status = tk.Label(rodape, text="", font=(FONTE, 9), fg=CINZA, bg=BRANCO, anchor="w")
lbl_status.pack(side="left", fill="x", expand=True)

barra_progresso = ttk.Progressbar(rodape, style="Preta.Horizontal.TProgressbar", length=180)

# --- Painel "Transações" ---
painel_transacoes = tk.Frame(janela, bg=BRANCO)

linha_info_t = tk.Frame(painel_transacoes, bg=BRANCO)
linha_info_t.pack(fill="x", pady=(0, 10))
lbl_arquivo = tk.Label(
    linha_info_t, text="Nenhuma planilha carregada", font=(FONTE, 10), fg=CINZA, bg=BRANCO, anchor="w"
)
lbl_arquivo.pack(side="left", fill="x", expand=True)

campo_busca = criar_campo(linha_info_t)
campo_busca.config(width=36)
campo_busca.insert(0, TEXTO_BUSCA)
campo_busca.config(fg=CINZA)
campo_busca.bind("<FocusIn>", busca_entrar)
campo_busca.bind("<FocusOut>", busca_sair)
campo_busca.bind("<KeyRelease>", agendar_busca)
lbl_lupa = tk.Label(linha_info_t, text="⌕", font=(FONTE, 14), fg=PRETO, bg=BRANCO)
# campo e lupa só aparecem depois de carregar a planilha

area_t, tabela_transacoes = criar_tabela(painel_transacoes, ("vazio",))
tabela_transacoes.heading("vazio", text="Clique em “Carregar transações” para abrir a planilha", anchor="w")
tabela_transacoes.column("vazio", width=600, stretch=True)
area_t.pack(fill="both", expand=True)

# --- Painel "Arquivos de comissão" ---
painel_comissoes = tk.Frame(janela, bg=BRANCO)

linha_info_c = tk.Frame(painel_comissoes, bg=BRANCO)
linha_info_c.pack(fill="x", pady=(0, 10))
lbl_pasta = tk.Label(
    linha_info_c, text="Nenhuma pasta selecionada", font=(FONTE, 10), fg=CINZA, bg=BRANCO, anchor="w"
)
lbl_pasta.pack(side="left", fill="x", expand=True)
tk.Label(
    linha_info_c, text="Duplo clique abre o arquivo", font=(FONTE, 9), fg=CINZA, bg=BRANCO
).pack(side="right")

colunas_c = ("gerar", "arquivo", "nome", "parceiro", "transacoes", "situacao")
area_c, tabela_comissoes = criar_tabela(painel_comissoes, colunas_c)
tabela_comissoes.heading("gerar", text="Gerar", anchor="center", command=alternar_todos)
tabela_comissoes.heading("arquivo", text="Nome do arquivo", anchor="w")
tabela_comissoes.heading("nome", text="Nome extraído", anchor="w")
tabela_comissoes.heading("parceiro", text="Parceiro", anchor="w")
tabela_comissoes.heading("transacoes", text="Transações", anchor="e")
tabela_comissoes.heading("situacao", text="Situação", anchor="w")
tabela_comissoes.column("gerar", width=64, minwidth=64, stretch=False, anchor="center")
tabela_comissoes.column("arquivo", width=380, minwidth=120, stretch=False, anchor="w")
tabela_comissoes.column("nome", width=170, minwidth=80, stretch=False, anchor="w")
tabela_comissoes.column("parceiro", width=220, minwidth=80, stretch=False, anchor="w")
tabela_comissoes.column("transacoes", width=110, minwidth=80, stretch=False, anchor="e")
tabela_comissoes.column("situacao", width=380, minwidth=300, stretch=True, anchor="w")
tabela_comissoes.bind("<Button-1>", clicar_na_tabela_comissoes)
tabela_comissoes.bind("<Double-1>", duplo_clique_comissoes)
area_c.pack(fill="both", expand=True)

mostrar_aba("transacoes")

# Ao voltar para o app (ex.: depois de editar o Arquivo - Parceiro), atualiza a tabela
janela.bind("<FocusIn>", recarregar_se_mapeamento_mudou)
janela.protocol("WM_DELETE_WINDOW", fechar_janela)
janela.after(80, verificar_fila)

janela.mainloop()
